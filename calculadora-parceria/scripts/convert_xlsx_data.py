"""Converte as abas de apoio de Cálculo_Parceria_REV02.xlsm em JSON para o app.

Uso:
    py scripts/convert_xlsx_data.py "C:/caminho/para/Cálculo_Parceria_REV02.xlsm"

Re-rode este script sempre que houver uma planilha atualizada (ex: novas
resoluções ANEEL) para atualizar os arquivos em src/data/.
"""
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

OUT_DIR = Path(__file__).resolve().parent.parent / "src" / "data"


def cell_value(ws, coord):
    v = ws[coord].value
    return v


def date_str(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return v


def convert_impostos(wb):
    ws = wb["IMPOSTOS"]
    rows = []
    for r in range(2, ws.max_row + 1):
        estado = ws.cell(r, 1).value
        if not estado:
            continue
        rows.append({
            "estado": estado,
            "icms": ws.cell(r, 2).value,
            "pisCofins": ws.cell(r, 3).value,
        })
    return rows


def convert_desconto(wb):
    ws = wb["DESCONTO"]
    rows = []
    for r in range(2, ws.max_row + 1):
        sigla = ws.cell(r, 1).value
        if not sigla:
            continue
        rows.append({
            "distribuidora": sigla,
            "estado": ws.cell(r, 2).value,
            "desconto": ws.cell(r, 3).value,
        })
    return rows


def convert_isencoes(wb):
    ws = wb["ISENCOES"]
    rows = []
    for r in range(2, ws.max_row + 1):
        estado = ws.cell(r, 1).value
        if not estado:
            continue
        rows.append({
            "estado": estado,
            "modalidade": ws.cell(r, 2).value,
            "fonte": ws.cell(r, 3).value,
            "faixa": ws.cell(r, 4).value,
            "teIcms": ws.cell(r, 5).value,
            "tePisCofins": ws.cell(r, 6).value,
            "tusdIcms": ws.cell(r, 7).value,
            "tusdPisCofins": ws.cell(r, 8).value,
        })
    return rows


def convert_equivalencia(wb):
    """Cada linha da planilha carrega dois apelidos para a mesma distribuidora:
    equiv1 (colunas A/B) é usado para casar com Tarifas_B1!C, equiv2 (colunas D/E)
    para casar com Componentes_Tarifarias!C. Na quase totalidade das linhas os dois
    apelidos coincidem; há pelo menos um caso real (Roraima Energia) onde divergem,
    então os dois precisam ser preservados separadamente por distribuidora."""
    ws = wb["Equivalencia"]
    rows = []
    for r in range(2, ws.max_row + 1):
        equiv1, nome1 = ws.cell(r, 1).value, ws.cell(r, 2).value
        equiv2, nome2 = ws.cell(r, 4).value, ws.cell(r, 5).value
        distribuidora = nome1 or nome2
        if not distribuidora:
            continue
        rows.append({
            "distribuidora": distribuidora,
            "equiv1": equiv1,
            "equiv2": equiv2,
        })
    return rows


def convert_tarifas_b1(wb):
    ws = wb["Tarifas_B1"]
    rows = []
    for r in range(2, ws.max_row + 1):
        codigo = ws.cell(r, 3).value
        if not codigo:
            continue
        rows.append({
            "codigo": codigo,
            "resolucao": ws.cell(r, 2).value,
            "inicioVigencia": date_str(ws.cell(r, 5).value),
            "fimVigencia": date_str(ws.cell(r, 6).value),
            "subgrupo": ws.cell(r, 8).value,
            "modalidade": ws.cell(r, 9).value,
            "classe": ws.cell(r, 10).value,
            "tusd": ws.cell(r, 16).value,
            "te": ws.cell(r, 17).value,
            "total": ws.cell(r, 18).value,
        })
    return rows


def convert_componentes_tarifarias(wb):
    ws = wb["Componentes_Tarifarias"]
    rows = []
    for r in range(2, ws.max_row + 1):
        codigo = ws.cell(r, 3).value
        if not codigo:
            continue
        rows.append({
            "codigo": codigo,
            "resolucao": ws.cell(r, 2).value,
            "inicioVigencia": date_str(ws.cell(r, 5).value),
            "fimVigencia": date_str(ws.cell(r, 6).value),
            "subgrupo": ws.cell(r, 8).value,
            "modalidade": ws.cell(r, 9).value,
            "classe": ws.cell(r, 10).value,
            "componente": ws.cell(r, 16).value,
            "valor": ws.cell(r, 17).value,
        })
    return rows


def convert_distribuidoras(wb):
    """Lista de distribuidoras (dropdown PARCERIA!C3 = BASE!$A$3:$A$105)."""
    ws = wb["BASE"]
    out = []
    for r in range(3, 106):
        v = ws.cell(r, 1).value
        if v:
            out.append(v)
    return out


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = Path(sys.argv[1])
    wb = load_workbook(src, data_only=True)

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    datasets = {
        "impostos.json": convert_impostos(wb),
        "desconto.json": convert_desconto(wb),
        "isencoes.json": convert_isencoes(wb),
        "equivalencia.json": convert_equivalencia(wb),
        "tarifas_b1.json": convert_tarifas_b1(wb),
        "componentes_tarifarias.json": convert_componentes_tarifarias(wb),
        "distribuidoras.json": convert_distribuidoras(wb),
    }

    for filename, data in datasets.items():
        path = OUT_DIR / filename
        path.write_text(json.dumps(data, ensure_ascii=False, indent=None), encoding="utf-8")
        print(f"{filename}: {len(data)} registros")


if __name__ == "__main__":
    main()
